"""
SITE BUILDER  -  bakes the latest run into the public Job Bot page
==================================================================

    python naukri_job_bot.py     # writes jobs.json
    python build_site.py         # writes site/index.html

The published page is a single self-contained file. It cannot fetch the
job data at view time (browsers block cross-site requests, and the
portals would block them anyway), so the data is embedded directly into
the HTML here, at build time.

That is also why "refreshing the site" means republishing it. The whole
refresh loop is:

    Task Scheduler -> naukri_job_bot.py -> build_site.py -> republish

Nothing about a visitor's resume is involved in any of this: resumes are
read and scored inside the visitor's own browser and never transmitted.
"""

import json
import os
import sys
from copy import deepcopy
from finance_catalog import public_catalog, classify_finance, listing_priority

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(SCRIPT_DIR, "site_template.html")
DATA = os.path.join(SCRIPT_DIR, "jobs.json")
OUT_DIR = os.path.join(SCRIPT_DIR, "site")
OUT = os.path.join(OUT_DIR, "index.html")
CONTACTS_XLSX = os.path.join(SCRIPT_DIR, "jobs.xlsx")

# The template carries this exact token where the payload belongs.
MARKER = '"__JOBS_PAYLOAD__"'


def _contact_key(source, title, company, location):
    """Stable join key shared by the public payload and local workbook."""
    def norm(value):
        return " ".join(str(value or "").lower().split())
    return tuple(norm(value) for value in (source, title, company, location))


def load_posting_contacts(path=CONTACTS_XLSX):
    """Read source-posted contacts from the local run workbook when present.

    Older jobs.json files intentionally contained no contacts. This bridge
    lets a rebuilt site use the already-collected workbook; new collections
    publish the same field directly through export_site_data.py.
    """
    if not path or not os.path.exists(path):
        return {}
    try:
        import openpyxl
        from export_site_data import publish_contact
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return {}
    contacts = {}
    for sheet_name in ("Shortlist", "Near Misses"):
        if sheet_name not in workbook.sheetnames:
            continue
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = next(rows, None)
        if not header or "Direct Contact" not in header:
            continue
        cols = {name: header.index(name) for name in
                ("Portal", "Role", "Company", "Location", "Direct Contact")}
        for row in rows:
            value = publish_contact(row[cols["Direct Contact"]])
            if value == "Not Available":
                continue
            key = _contact_key(row[cols["Portal"]], row[cols["Role"]],
                               row[cols["Company"]], row[cols["Location"]])
            contacts[key] = value
    workbook.close()
    return contacts


def load_payload(path=DATA):
    if not os.path.exists(path):
        print(f"  ! {os.path.basename(path)} not found.")
        print("    Run 'python naukri_job_bot.py' first - it writes jobs.json")
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def build(payload, template_path=TEMPLATE, out_path=OUT, contacts_path=CONTACTS_XLSX):
    with open(template_path, encoding="utf-8") as f:
        html = f.read()

    if MARKER not in html:
        raise SystemExit(f"Template is missing the {MARKER} marker.")

    # Reclassify existing public rows too; new searches populate on collection.
    # A listing score is independent of the collector owner's salary/experience.
    payload = deepcopy(payload)
    posting_contacts = load_posting_contacts(contacts_path)
    from export_site_data import redact_contacts
    for job in payload.get("jobs", []):
        family, basis, strength = classify_finance(job)
        if family and (basis == "title" or strength >= .3):
            job["family"] = family
        job["score"] = listing_priority(job)
        job["why"] = "Listing priority uses posting freshness, employer verification and reported applicant counts. It does not predict interview chances."
        if job.get("salary_basis", "").startswith(("Est.", "Not disclosed", "Unknown")):
            job.update(salary_basis="Not disclosed", salary_band="-", sal_lo=None, sal_hi=None)
        # Backfill older payloads so the application-method filter is useful
        # immediately after a template-only rebuild. LinkedIn is deliberately
        # not guessed: Easy Apply is populated only by collection evidence.
        if job.get("source") == "Workday" and not job.get("apply_method"):
            job["apply_method"] = "Employer site"
        key = _contact_key(job.get("source"), job.get("title"),
                           job.get("company"), job.get("location"))
        job["contact"] = posting_contacts.get(key, job.get("contact"))
        if job.get("contact") in (None, "", "-"):
            job["contact"] = "Not Available"
        # Older public payloads may predate the stricter punctuation-aware
        # phone redaction. Keep contact details in their labelled field only.
        for field in ("title", "company", "why", "missing", "skills", "jd", "requirement_text"):
            job[field] = redact_contacts(job.get(field, ""))
    payload["finance_catalog"] = public_catalog()
    payload.setdefault("counts", {})["by_family"] = {}
    for job in payload.get("jobs", []):
        family = job.get("family", "Finance")
        payload["counts"]["by_family"][family] = payload["counts"]["by_family"].get(family, 0) + 1

    with open(os.path.join(SCRIPT_DIR, "finance_match.js"), encoding="utf-8") as f:
        html = html.replace("/*__FINANCE_MATCH__*/", f.read())

    # json.dumps output is safe to drop into a <script type="application/json">
    # block except for a literal "</script>" inside a string value, which
    # would close the tag early. Escaping the slash keeps the JSON valid.
    blob = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    blob = blob.replace("</", "<\\/")

    html = html.replace(MARKER, blob)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


def main():
    payload = load_payload()
    if payload is None:
        return 1

    out = build(payload)
    counts = payload.get("counts", {})
    by_source = counts.get("by_source", {})
    size_kb = os.path.getsize(out) // 1024

    print("Job Bot site built")
    print(f"  data from : {payload.get('generated_label', '?')}")
    print(f"  jobs      : {counts.get('total', 0)} "
          f"({counts.get('shortlist', 0)} shortlist, "
          f"{counts.get('near', 0)} near-miss)")
    for src, n in sorted(by_source.items(), key=lambda kv: -kv[1]):
        print(f"      {src:<10}: {n}")
    print(f"  matchable : {counts.get('with_text', 0)} jobs carry JD text")
    print(f"  -> {out} ({size_kb} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
